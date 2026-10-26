"""Jarvis - Registon maktabi marketing yordamchisi (Telegram bot).

Barcha maxfiy qiymatlar faqat muhit o'zgaruvchilaridan (env) o'qiladi:
  TELEGRAM_BOT_TOKEN, ANTHROPIC_API_KEY, ALLOWED_USER_IDS, CLAUDE_MODEL
  WEBHOOK_URL (ixtiyoriy; Render'da RENDER_EXTERNAL_URL avtomatik ishlatiladi)
  GROQ_API_KEY (ixtiyoriy; ovozli xabarlarni matnga aylantirish uchun)
"""
import hashlib
import io
import logging
import os
import zipfile
from collections import defaultdict, deque

import httpx
from anthropic import AsyncAnthropic
from telegram import Update
from telegram.constants import ChatAction
from telegram.ext import Application, CommandHandler, ContextTypes, MessageHandler, filters

logging.basicConfig(format="%(asctime)s %(levelname)s %(message)s", level=logging.INFO)
logging.getLogger("httpx").setLevel(logging.WARNING)  # so'rov URL'larida token bor, logga chiqmasin
log = logging.getLogger("jarvis")

def _env(name: str) -> str:
    value = os.environ.get(name, "").strip().strip('"').strip("'")
    if not value:
        raise SystemExit(f"XATO: {name} muhit o'zgaruvchisi kiritilmagan (Render -> Environment).")
    return value


TOKEN = _env("TELEGRAM_BOT_TOKEN")
os.environ["ANTHROPIC_API_KEY"] = _env("ANTHROPIC_API_KEY")
MODEL = os.environ.get("CLAUDE_MODEL", "claude-sonnet-5-5")
ALLOWED = {int(x) for x in os.environ.get("ALLOWED_USER_IDS", "").replace(" ", "").split(",") if x}
GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "").strip()
HISTORY_LEN = 20  # har bir chat uchun eslab qolinadigan oxirgi xabarlar soni

SYSTEM_PROMPT = """Sen Jarvis - Shodiyorning marketing yordamchisisan.
Shodiyor Registon maktabida (xususiy maktab) marketing bilan shug'ullanadi va marketing konsalting qiladi.
Targetolog sherigi - Muzaffar. Qabul mavsumlari: 1-mavsum 1200 o'quvchi, rekord mavsum 1800 o'quvchi.
Maktab dushanbadan shanbagacha ishlaydi.

Qoidalar:
- Har doim lotin yozuvidagi o'zbek tilida javob ber, "Shodiyor" deb murojaat qil.
- Qisqa va aniq: avval natija, keyin izoh. Maqtov uchun maqtama, xato ko'rsang ochiq ayt.
- Raqam o'ylab topma: ma'lumot bo'lmasa so'ra yoki "taxminiy" deb belgila.
- Ota-onalarni chalg'itadigan yoki tekshirib bo'lmaydigan da'volar yozma ("eng yaxshi maktab" kabi bo'sh va'dalar emas, aniq faktlar).

Video ssenariy: platforma, davomiylik, maqsad, kim gapirishini aniqla (aytilmasa o'zing taklif qil).
Ssenariyni ustunlar bilan ber: vaqt | kadrda nima | ovoz matni | ekrandagi yozuv.
Birinchi 3 soniyada ota-onaning real tashvishidan ilmoq; oxirida aniq chaqiruv (qo'ng'iroq, Telegram bot, ro'yxatdan o'tish).
So'ralmasa ham 2-3 xil ilmoq varianti taklif qil.

Post va reklama matni: 3 variant - muammo-yechim, natija va raqamlar, ota-ona hikoyasi. Birinchi qator - ilmoq.

Tahlil: CPL, lid -> qabul konversiyasi, CAC, kanallar bo'yicha taqqoslash.
Tuzilma: asosiy xulosa 2-3 gapda; nima ishlayapti va nima ishlamayapti; keyingi haftaga 3 ta aniq harakat.

Telegram'da o'qish qulay bo'lsin: jadval o'rniga raqamlangan bandlar ishlat, markdown belgilarini (**, ##) ishlatma.

Shodiyor yuklagan fayllar suhbat boshida beriladi. Savol ularga tegishli bo'lsa, javobni shu fayllardan ol
va qaysi fayldan olganingni ayt. Faylda javob bo'lmasa, "fayllarda topilmadi" deb ochiq ayt.
Faqat so'ralgan aniq javobni ber (raqam, narx, ism, sana yoki kerakli band). Faylni yoki uning katta qismini
ko'chirib berma, "faylga qarang" dema: o'zing topib, qisqa javob ber."""

client = AsyncAnthropic()  # ANTHROPIC_API_KEY env'dan o'qiladi
history: dict[int, deque] = defaultdict(lambda: deque(maxlen=HISTORY_LEN))

# Fayllar Anthropic Files API'da saqlanadi: Render'ning bepul diski server uxlaganda tozalanadi,
# Files API'dagi fayllar esa o'chirilmaguncha turadi. Bot fayllari shu prefiks bilan ajratiladi.
FILE_PREFIX = "jarvis__"
KB_PREFIX = FILE_PREFIX + "kb_"  # repo'dagi knowledge/ papkasidan kelgan fayllar
KB_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "knowledge")
_files_cache: list | None = None


async def sync_knowledge(files: list) -> list:
    """knowledge/ papkasini Files API bilan moslaydi: yangi fayllarni yuklaydi, olib tashlanganlarini o'chiradi.

    Fayl nomiga mazmun xeshi qo'shiladi, shuning uchun o'zgargan fayl qayta yuklanadi."""
    wanted = {}
    if os.path.isdir(KB_DIR):
        for name in sorted(os.listdir(KB_DIR)):
            path = os.path.join(KB_DIR, name)
            if not os.path.isfile(path) or name.startswith("."):
                continue
            with open(path, "rb") as fh:
                data = fh.read()
            prepared = to_upload(name, data)
            if prepared is None:
                log.warning("knowledge/: o'qib bo'lmadi, o'tkazildi: %s", name)
                continue
            up_name, content, mime = prepared
            digest = hashlib.sha256(content).hexdigest()[:8]
            wanted[f"{KB_PREFIX}{digest}_{up_name}"] = (content, mime)

    kept = []
    for f in files:
        if f.filename.startswith(KB_PREFIX) and f.filename not in wanted:
            await client.files.delete(f.id)
            log.info("knowledge/: o'chirildi %s", f.filename)
        else:
            kept.append(f)
    have = {f.filename for f in kept}
    for fname, (content, mime) in wanted.items():
        if fname not in have:
            kept.append(await client.files.upload(file=(fname, content, mime)))
            log.info("knowledge/: yuklandi %s", fname)
    return kept


def display_name(f) -> str:
    name = f.filename.removeprefix(FILE_PREFIX)
    if f.filename.startswith(KB_PREFIX):
        name = name.removeprefix("kb_").split("_", 1)[-1] + " (doimiy)"
    return name


async def list_files() -> list:
    global _files_cache
    if _files_cache is None:
        files = [f async for f in client.files.list() if f.filename.startswith(FILE_PREFIX)]
        files = await sync_knowledge(files)
        files.sort(key=lambda f: f.created_at)
        _files_cache = files
    return _files_cache


def to_upload(name: str, data: bytes) -> tuple[str, bytes, str] | None:
    """Telegram faylini Files API uchun (nom, baytlar, mime) ko'rinishiga keltiradi."""
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext == "pdf":
        return name, data, "application/pdf"
    if ext in ("txt", "md", "csv"):
        text = data.decode("utf-8", errors="replace")
    elif ext == "docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            parts += ["\t".join(c.text for c in row.cells) for row in table.rows]
        text = "\n".join(parts)
    elif ext == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"## {ws.title}")
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    parts.append("\t".join("" if v is None else str(v) for v in row))
        text = "\n".join(parts)
    else:
        return None
    return name.rsplit(".", 1)[0] + ".txt", text.encode("utf-8"), "text/plain"


async def allowed(update: Update) -> bool:
    uid = update.effective_user.id
    if not ALLOWED:
        await update.message.reply_text(
            f"Sizning Telegram ID: {uid}\nUni ALLOWED_USER_IDS ga yozing va botni qayta ishga tushiring."
        )
        return False
    if uid not in ALLOWED:
        log.warning("Ruxsatsiz urinish: user_id=%s", uid)
        await update.message.reply_text("Kechirasiz, bu bot shaxsiy.")
        return False
    return True


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if await allowed(update):
        await update.message.reply_text(
            "Salom, Shodiyor! Men Jarvisman. Ssenariy, post, reklama matni yoki tahlil - nima kerak?\n"
            "/new - suhbatni yangidan boshlash\n"
            "Fayl yoki ZIP arxiv yuborsangiz (PDF, Word, Excel, TXT), saqlab qo'yaman va savollaringizga undan javob beraman.\n"
            "/fayllar - saqlangan fayllar ro'yxati\n"
            "/ochir 2 - 2-faylni o'chirish\n"
            "Ovozli xabar ham yuborsangiz bo'ladi."
        )


async def new_chat(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if await allowed(update):
        history.pop(update.effective_chat.id, None)
        await update.message.reply_text("Yangi suhbat boshlandi.")


async def save_document(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await allowed(update):
        return
    doc = update.message.document
    if doc.file_size and doc.file_size > 20 * 1024 * 1024:
        await update.message.reply_text("Fayl 20 MB dan katta, Telegram botlari buni yuklab ololmaydi.")
        return
    tg_file = await doc.get_file()
    data = bytes(await tg_file.download_as_bytearray())
    name = doc.file_name or "fayl"

    if name.lower().endswith(".zip"):
        # Papkani ZIP qilib yuborish mumkin: ichidagi har bir fayl alohida saqlanadi
        items = []
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            for info in zf.infolist():
                base = info.filename.rsplit("/", 1)[-1]
                if info.is_dir() or not base or base.startswith(("~$", ".")) or "__MACOSX" in info.filename:
                    continue
                items.append((base, zf.read(info)))
    else:
        items = [(name, data)]

    saved, skipped = [], []
    for item_name, item_data in items:
        try:
            prepared = to_upload(item_name, item_data)
        except Exception:
            log.exception("Faylni o'qib bo'lmadi: %s", item_name)
            prepared = None
        if prepared is None:
            skipped.append(item_name)
            continue
        up_name, content, mime = prepared
        meta = await client.files.upload(file=(FILE_PREFIX + up_name, content, mime))
        (await list_files()).append(meta)
        saved.append(up_name)

    lines = []
    if saved:
        lines.append(f"Saqlandi ({len(saved)}): " + ", ".join(saved))
    if skipped:
        lines.append(f"O'qib bo'lmadi ({len(skipped)}): " + ", ".join(skipped))
        lines.append("Qabul qilinadigan turlar: PDF, DOCX, XLSX, TXT, CSV, MD.")
    if saved:
        lines.append("Endi shu fayllar bo'yicha savol bersangiz bo'ladi.")
    await update.message.reply_text("\n".join(lines) or "ZIP ichida fayl topilmadi.")


async def show_files(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await allowed(update):
        return
    files = await list_files()
    if not files:
        await update.message.reply_text("Hali fayl yo'q. Menga PDF, Word yoki Excel fayl yuboring.")
        return
    lines = [f"{i}. {display_name(f)}" for i, f in enumerate(files, 1)]
    await update.message.reply_text("Saqlangan fayllar:\n" + "\n".join(lines))


async def delete_file(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await allowed(update):
        return
    files = await list_files()
    try:
        idx = int(context.args[0]) - 1
        assert 0 <= idx < len(files)
    except (IndexError, ValueError, AssertionError):
        await update.message.reply_text("Raqamni yozing, masalan: /ochir 2  (ro'yxat: /fayllar)")
        return
    if files[idx].filename.startswith(KB_PREFIX):
        await update.message.reply_text("Bu doimiy fayl (GitHub'dagi knowledge papkasidan). Uni o'sha yerdan olib tashlash kerak.")
        return
    f = files.pop(idx)
    await client.files.delete(f.id)
    await update.message.reply_text(f"O'chirildi: {display_name(f)}")


async def files_prefix() -> list:
    """Fayllarni suhbat boshiga qo'yadi; kesh belgisi qayta yuborishni arzonlashtiradi."""
    files = await list_files()
    if not files:
        return []
    blocks = [
        {"type": "document", "source": {"type": "file", "file_id": f.id},
         "title": display_name(f)}
        for f in files
    ]
    blocks[-1]["cache_control"] = {"type": "ephemeral"}
    blocks.append({"type": "text", "text": "Bular men yuklagan fayllar. Kerak bo'lganda shulardan foydalan."})
    return [
        {"role": "user", "content": blocks},
        {"role": "assistant", "content": "Fayllarni ko'rdim, savolingizni kutaman."},
    ]


async def transcribe(audio: bytes) -> str:
    """Ovozni Groq Whisper orqali matnga aylantiradi."""
    async with httpx.AsyncClient(timeout=60) as http:
        r = await http.post(
            "https://api.groq.com/openai/v1/audio/transcriptions",
            headers={"Authorization": f"Bearer {GROQ_API_KEY}"},
            files={"file": ("voice.ogg", audio, "audio/ogg")},
            data={"model": "whisper-large-v3", "language": "uz", "response_format": "json"},
        )
        r.raise_for_status()
        return r.json().get("text", "").strip()


async def handle_voice(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await allowed(update):
        return
    if not GROQ_API_KEY:
        await update.message.reply_text("Ovozni o'qish uchun Render'ga GROQ_API_KEY qo'shilishi kerak.")
        return
    voice = update.message.voice or update.message.audio
    await context.bot.send_chat_action(update.effective_chat.id, ChatAction.TYPING)
    try:
        audio = bytes(await (await voice.get_file()).download_as_bytearray())
        text = await transcribe(audio)
    except Exception:
        log.exception("Ovozni matnga aylantirib bo'lmadi")
        await update.message.reply_text("Ovozni tushuna olmadim, qayta yuboring yoki matn bilan yozing.")
        return
    if not text:
        await update.message.reply_text("Ovozda gap eshitilmadi.")
        return
    await update.message.reply_text(f"Eshitdim: {text}")
    await answer(update, context, text)


async def handle(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if await allowed(update):
        await answer(update, context, update.message.text)


async def answer(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str) -> None:
    chat_id = update.effective_chat.id
    msgs = history[chat_id]
    msgs.append({"role": "user", "content": text})
    await context.bot.send_chat_action(chat_id, ChatAction.TYPING)

    # Tarix deque bilan qirqilganda birinchi xabar "user" bo'lishi shart
    convo = list(msgs)
    while convo and convo[0]["role"] != "user":
        convo.pop(0)

    try:
        resp = await client.messages.create(
            model=MODEL, max_tokens=4000, system=SYSTEM_PROMPT, messages=await files_prefix() + convo
        )
        answer = "".join(b.text for b in resp.content if b.type == "text").strip() or "(bo'sh javob)"
    except Exception:
        log.exception("Claude API xatosi")
        msgs.pop()
        await update.message.reply_text("Xatolik yuz berdi, birozdan keyin qayta urinib ko'ring.")
        return

    msgs.append({"role": "assistant", "content": answer})
    for i in range(0, len(answer), 4000):  # Telegram cheklovi 4096 belgi
        await update.message.reply_text(answer[i : i + 4000])


def main() -> None:
    app = Application.builder().token(TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("new", new_chat))
    app.add_handler(CommandHandler("fayllar", show_files))
    app.add_handler(CommandHandler("ochir", delete_file))
    app.add_handler(MessageHandler(filters.Document.ALL, save_document))
    app.add_handler(MessageHandler(filters.VOICE | filters.AUDIO, handle_voice))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle))
    log.info("Jarvis ishga tushdi (model: %s)", MODEL)

    # Render bepul Web Service'da RENDER_EXTERNAL_URL avtomatik beriladi -> webhook rejimi.
    # Server uxlab qolsa, Telegram'dan kelgan xabar uni uyg'otadi.
    public_url = os.environ.get("WEBHOOK_URL") or os.environ.get("RENDER_EXTERNAL_URL")
    if public_url:
        path = hashlib.sha256(TOKEN.encode()).hexdigest()[:32]  # tokenni URL'da ochmaslik uchun
        app.run_webhook(
            listen="0.0.0.0",
            port=int(os.environ.get("PORT", "10000")),
            url_path=path,
            webhook_url=f"{public_url.rstrip('/')}/{path}",
            secret_token=path,
            allowed_updates=Update.ALL_TYPES,
        )
    else:
        app.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
